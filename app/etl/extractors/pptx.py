"""python-pptx extractor for PowerPoint files (shapes, tables, notes)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from app.etl.extractors.base import ExtractedDocument, ExtractorError


def _shape_text(shape) -> str:
    texts: list[str] = []
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            run_text = "".join(run.text for run in para.runs)
            if run_text.strip():
                texts.append(run_text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            texts.append(" | ".join(cells))
    if shape.shape_type == 6:  # group shape
        for sub in shape.shapes:
            sub_text = _shape_text(sub)
            if sub_text.strip():
                texts.append(sub_text)
    return "\n".join(texts)


def extract_pptx(path: Path) -> ExtractedDocument:
    try:
        prs = Presentation(str(path))
    except Exception as exc:
        raise ExtractorError(f"python-pptx failed to open presentation: {exc}") from exc

    slides: list[dict] = []
    for idx, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            text = _shape_text(shape)
            if text.strip():
                parts.append(text)
        # Speaker notes are often the densest text in slide decks.
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[notes]\n{notes}")
        slide_text = "\n\n".join(parts)
        slides.append({"page": idx, "text": slide_text})

    text = "\n\n".join(s.get("text", "") for s in slides if s.get("text"))
    table_count = sum(
        sum(1 for shape in slide.shapes if getattr(shape, "has_table", False) and shape.has_table)
        for slide in prs.slides
    )

    warnings: list[str] = []
    if not text.strip():
        warnings.append("No text extracted from presentation (image-only slides?)")
    return ExtractedDocument(
        text=text,
        pages=slides,
        page_count=len(prs.slides),
        warnings=warnings,
        confidence=None,
        ocr_used=False,
        tables=table_count,
    )