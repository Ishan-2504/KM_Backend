"""python-docx extractor for Word documents (paragraphs + tables)."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from app.etl.extractors.base import ExtractedDocument, ExtractorError


def _table_text(table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        rows.append(" | ".join(cells))
    return "\n".join(rows)


def extract_docx(path: Path) -> ExtractedDocument:
    try:
        doc = Document(str(path))
    except Exception as exc:
        raise ExtractorError(f"python-docx failed to open document: {exc}") from exc

    # Walk the document body so paragraphs and tables keep their flow order.
    paragraphs = iter(doc.paragraphs)
    parts: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = next(paragraphs, None)
            if para is not None and para.text.strip():
                parts.append(para.text)
        elif tag == "tbl":
            for table in doc.tables:
                if table._tbl is child:
                    parts.append(_table_text(table))
                    break

    text = "\n\n".join(p for p in parts if p.strip())
    warnings: list[str] = []
    if not text.strip():
        warnings.append("No paragraph or table text extracted")
    return ExtractedDocument(
        text=text,
        pages=[],
        page_count=1,
        warnings=warnings,
        confidence=None,
        ocr_used=False,
        tables=len(doc.tables),
    )