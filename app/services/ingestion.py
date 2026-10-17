"""Offline ingestion: embed documents and write the FAISS index.

Two sources are supported:
  - Excel workbook (`DATA_FILE_PATH`), the legacy metadata-driven source.
  - ETL JSONL (`ETL_RECORDS_FILE`), canonical records produced by
    `python -m app.etl.pipeline`. Metadata fields are populated from the
    record's `metadata` block (empty for now - populated in a later phase),
    and the extracted text drives the composite text for document-level
    embedding.
"""

import argparse
import hashlib
import sys
from pathlib import Path

from openpyxl import load_workbook
from tqdm import tqdm

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from app.config import DATA_FILE_PATH, ETL_RECORDS_FILE, INDEX_MAX_EMBED_CHARS, VECTOR_DIMENSIONS
from app.db.search_client import ensure_index, get_search_client
from app.etl.models import ETLRecord
from app.models import COLUMN_MAP, DOC_TYPE_MAP, METADATA_FIELDS, NULL_MARKERS
from app.services.embeddings import BATCH_SIZE, embed_texts

UPLOAD_CHUNK = 100


def _resolve_path(value: str | None, default: str) -> Path:
    path = Path(value or default)
    if not path.is_absolute():
        path = Path(__file__).resolve().parent.parent.parent / path
    return path


def _normalise_value(value) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    if s in NULL_MARKERS:
        return None
    return s


def _doc_id(name: str, idx: int) -> str:
    return hashlib.md5(f"{name}_{idx}".encode()).hexdigest()


# ---------------------------------------------------------------- Excel source

def load_records(file_path: str | None = None) -> list[dict]:
    path = _resolve_path(file_path, DATA_FILE_PATH)
    wb = load_workbook(str(path), read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    headers = rows[0]
    records: list[dict] = []
    for row in rows[1:]:
        raw = {str(headers[i]): row[i] for i in range(len(headers))}
        record: dict = {}
        for src_col, dst_col in COLUMN_MAP.items():
            record[dst_col] = _normalise_value(raw.get(src_col))

        if record.get("document_type") in DOC_TYPE_MAP:
            record["document_type"] = DOC_TYPE_MAP[record["document_type"]]

        records.append(record)

    return records


def _build_composite_text(rec: dict) -> str:
    """Build a search-optimised composite string for embedding.

    Bios embed name, designation, and summary so role / skill queries match.
    KL docs embed name, industry, knowledge type, solution, and summary
    so structured metadata signals are captured in the vector.
    """
    doc_type = rec.get("document_type", "")
    if doc_type == "bio":
        parts = [
            rec.get("document_name", ""),
            rec.get("designation", ""),
            rec.get("summary", ""),
        ]
    else:
        parts = [
            rec.get("document_name", ""),
            rec.get("industry", ""),
            rec.get("sub_industry", ""),
            rec.get("knowledge_type", ""),
            rec.get("solution", ""),
            rec.get("summary", ""),
        ]
    return " | ".join(p for p in parts if p)


# ---------------------------------------------------------------- ETL source

def load_etl_records(records_path: str | None = None) -> list[dict]:
    """Read canonical ETL records (JSONL) into ingestion-shaped rows."""
    path = _resolve_path(records_path, ETL_RECORDS_FILE)
    if not path.exists():
        raise FileNotFoundError(
            f"ETL records file not found: {path}. Run the ETL pipeline first "
            f"(python -m app.etl.pipeline --out {path.parent})."
        )

    records: list[dict] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        rec = ETLRecord.model_validate_json(line)
        row: dict = {}
        for field in METADATA_FIELDS:
            row[field] = (rec.metadata.get(field) or "").strip()
        row["doc_id"] = rec.doc_id
        row["document_name"] = rec.document_name
        row["summary"] = rec.summary
        row["file_text"] = rec.text
        records.append(row)
    return records


def _build_etl_composite_text(rec: dict) -> str:
    """Document-level composite for ETL records: name + metadata signals +
    a bounded slice of the extracted text (metadata is empty for now)."""
    name = rec.get("document_name") or ""
    meta = " | ".join(
        p for p in (rec.get(f) or "" for f in (
            "industry", "sub_industry", "knowledge_type", "solution", "client_name",
        ))
        if p
    )
    snippet = (rec.get("file_text") or "").strip()
    if len(snippet) > INDEX_MAX_EMBED_CHARS:
        snippet = snippet[:INDEX_MAX_EMBED_CHARS] + " ..."
    return "\n".join(p for p in (name, meta, snippet) if p)


def _build_documents(records: list[dict], etl: bool) -> tuple[list[dict], list[str]]:
    documents: list[dict] = []
    texts_to_embed: list[str] = []
    for idx, rec in enumerate(records):
        if etl:
            doc_id = rec.get("doc_id") or _doc_id(rec.get("document_name", ""), idx)
            composite = _build_etl_composite_text(rec)
        else:
            doc_id = _doc_id(rec.get("document_name", ""), idx)
            composite = _build_composite_text(rec)

        summary = rec.get("summary") or rec.get("document_name", "")
        doc = {
            "id": doc_id,
            "document_name": rec.get("document_name") or "",
            "summary": summary,
            "composite_text": composite,
            "file_text": rec.get("file_text") or "",
        }
        for field in METADATA_FIELDS:
            doc[field] = rec.get(field) or ""

        documents.append(doc)
        texts_to_embed.append(composite)
    return documents, texts_to_embed


def _embed_and_upload(client, documents: list[dict], texts_to_embed: list[str]) -> int:
    if not documents:
        print("No documents to index.")
        return 0

    print(f"Embedding {len(texts_to_embed)} documents in batches of {BATCH_SIZE}...")
    all_embeddings: list[list[float]] = []
    for i in tqdm(range(0, len(texts_to_embed), BATCH_SIZE), desc="Embedding"):
        batch = texts_to_embed[i : i + BATCH_SIZE]
        all_embeddings.extend(embed_texts(batch))

    for i, doc in enumerate(documents):
        doc["embedding"] = all_embeddings[i]

    print(f"Uploading {len(documents)} documents to FAISS...")
    for i in tqdm(range(0, len(documents), UPLOAD_CHUNK), desc="Uploading"):
        client.upload_documents(documents[i : i + UPLOAD_CHUNK])

    count = client.get_document_count()
    print(f"Ingestion complete. Index now has {count} documents.")
    return count


# ---------------------------------------------------------------- entry points

def ingest(file_path: str | None = None) -> int:
    """Index from the Excel workbook (legacy source)."""
    ensure_index()
    records = load_records(file_path)
    client = get_search_client()
    client.reset()
    documents, texts_to_embed = _build_documents(records, etl=False)
    return _embed_and_upload(client, documents, texts_to_embed)


def ingest_etl(records_path: str | None = None) -> int:
    """Index from canonical ETL records (JSONL)."""
    ensure_index()
    records = load_etl_records(records_path)
    client = get_search_client()
    client.reset()
    documents, texts_to_embed = _build_documents(records, etl=True)
    return _embed_and_upload(client, documents, texts_to_embed)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(prog="ingestion", description="Build the FAISS index.")
    parser.add_argument(
        "--etl", metavar="RECORDS_JSONL", default=None,
        help="index from ETL records.jsonl instead of the Excel workbook",
    )
    args = parser.parse_args()
    if args.etl:
        ingest_etl(args.etl)
    else:
        ingest()